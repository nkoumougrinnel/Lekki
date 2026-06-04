"""
Lekki Wiki — ImportService

Pipeline d'import documentaire intelligent :
  fichier (PDF/DOCX/TXT/MD) → extraction texte → création Page (workspace courant)
  → chunking + vectorisation + embeddings (RAG) → page consultable par le chatbot.

Le traitement lourd s'exécute en tâche de fond (FastAPI BackgroundTasks) sur une
session de base de données dédiée, car la session de la requête est fermée dès la
réponse envoyée.
"""

from __future__ import annotations

import io
import os
from datetime import UTC, datetime

from sqlalchemy import select, text

from app.database import AsyncSessionLocal
from app.models.import_job import Import
from app.models.page import Page
from app.services import rag_service

# Extensions supportées → type de source normalisé.
SUPPORTED_TYPES: dict[str, str] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".txt": "txt",
    ".md": "markdown",
    ".markdown": "markdown",
}


def detect_source_type(filename: str) -> str | None:
    """Retourne le type de source pour un nom de fichier, ou None si non supporté."""
    _, ext = os.path.splitext(filename.lower())
    return SUPPORTED_TYPES.get(ext)


def is_supported(filename: str) -> bool:
    return detect_source_type(filename) is not None


def title_from_filename(filename: str) -> str:
    """Dérive un titre lisible depuis le nom de fichier."""
    base = os.path.basename(filename)
    name, _ = os.path.splitext(base)
    cleaned = name.replace("_", " ").replace("-", " ").strip()
    return cleaned or "Document importé"


# ---------------------------------------------------------------------------
# Extraction de texte
# ---------------------------------------------------------------------------

def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n\n".join(p.strip() for p in parts if p.strip())


def _extract_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    # Inclure le texte des tableaux.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _extract_plain(data: bytes) -> str:
    return data.decode("utf-8", errors="replace").strip()


def extract_text(filename: str, data: bytes) -> str:
    """Extrait le texte d'un fichier selon son extension."""
    source_type = detect_source_type(filename)
    if source_type == "pdf":
        return _extract_pdf(data)
    if source_type == "docx":
        return _extract_docx(data)
    if source_type in ("txt", "markdown"):
        return _extract_plain(data)
    raise ValueError(f"Type de fichier non supporté : {filename}")


# ---------------------------------------------------------------------------
# Pipeline de traitement (tâche de fond)
# ---------------------------------------------------------------------------

async def _sync_fts(session, page: Page) -> None:
    try:
        await session.execute(
            text("DELETE FROM pages_fts WHERE page_id = :id"), {"id": page.id}
        )
        await session.execute(
            text(
                "INSERT INTO pages_fts (page_id, title, content) "
                "VALUES (:id, :title, :content)"
            ),
            {"id": page.id, "title": page.title, "content": page.content},
        )
        await session.commit()
    except Exception as exc:  # FTS non bloquant
        print(f"[import] FTS sync error: {exc}")


async def process_import(
    import_id: str,
    files: list[tuple[str, bytes]],
    user_id: str,
    workspace_id: str | None,
    category: str = "guides",
) -> None:
    """
    Traite un lot de fichiers en arrière-plan :
    extraction → Page → RAG (chunking + embeddings). Met à jour la progression.
    """
    async with AsyncSessionLocal() as session:
        job = await session.get(Import, import_id)
        if not job:
            return

        job.status = "processing"
        await session.commit()

        errors: list[str] = []
        processed = 0

        for filename, data in files:
            try:
                content_text = extract_text(filename, data)
                if not content_text.strip():
                    raise ValueError("aucun texte extractible")

                title = title_from_filename(filename)
                # Préfixer un titre Markdown si absent (meilleur rendu + RAG).
                body = content_text
                if not body.lstrip().startswith("#"):
                    body = f"# {title}\n\n{content_text}"

                page = Page(
                    title=title,
                    content=body,
                    category=category,
                    status="published",
                    creator_id=user_id,
                    workspace_id=workspace_id,
                )
                session.add(page)
                await session.commit()
                await session.refresh(page)

                await _sync_fts(session, page)

                # Chunking + vectorisation + stockage des embeddings.
                await rag_service.embed_page(session, page.id)

                processed += 1
                job.processed_files = processed
                await session.commit()
            except Exception as exc:
                errors.append(f"{filename}: {exc}")
                # Repartir d'une session propre après une erreur éventuelle.
                await session.rollback()
                job = await session.get(Import, import_id)
                job.processed_files = processed
                job.error_log = "\n".join(errors)
                await session.commit()

        # Statut final.
        job = await session.get(Import, import_id)
        job.processed_files = processed
        job.error_log = "\n".join(errors) if errors else None
        if processed == 0:
            job.status = "failed"
        elif errors:
            job.status = "partial"
        else:
            job.status = "completed"
        await session.commit()

        # Recalcul des pages liées si au moins une page a été créée.
        if processed > 0:
            try:
                await rag_service.compute_related_pages(session)
            except Exception as exc:
                print(f"[import] calcul des pages liées échoué: {exc}")


async def list_user_imports(user_id: str) -> list[Import]:
    async with AsyncSessionLocal() as session:
        result = await session.execute(
            select(Import)
            .where(Import.user_id == user_id)
            .order_by(Import.created_at.desc())
        )
        return list(result.scalars().all())
